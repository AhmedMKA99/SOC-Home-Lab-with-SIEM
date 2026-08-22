from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

for section in doc.sections:
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

NAVY = (0x2E, 0x40, 0x57)
GREY = (110, 110, 110)


def set_font(run, bold=False, size=11, color=None, italic=False):
    run.font.name = 'Calibri'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def hr(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '2E4057')
    pBdr.append(bottom)
    pPr.append(pBdr)


def h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    set_font(p.add_run(text.upper()), bold=True, size=13, color=NAVY)
    hr(doc)


def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run(text), bold=True, size=12, color=NAVY)


def body(doc, text, italic=False, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    set_font(p.add_run(text), size=11, italic=italic, color=color)


def bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    set_font(p.add_run(text), size=11)


def kv_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = 'Light Shading Accent 1'
    for i, (k, v) in enumerate(rows):
        table.rows[i].cells[0].text = k
        table.rows[i].cells[1].text = v
        for cell in table.rows[i].cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.name = 'Calibri'
                    run.font.size = Pt(10.5)


def timeline_table(doc, rows):
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Medium Shading 1 Accent 1'
    hdr = table.rows[0].cells
    hdr[0].text = 'Timestamp (UTC)'
    hdr[1].text = 'Event'
    hdr[2].text = 'Source'
    for k in hdr:
        for p in k.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(10)
    for ts, ev, src in rows:
        row = table.add_row().cells
        row[0].text = ts
        row[1].text = ev
        row[2].text = src
        for cell in row:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)


def alert_table(doc, rows):
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Medium Shading 1 Accent 1'
    hdr = table.rows[0].cells
    for i, h in enumerate(['Rule Description', 'Approx. Rule ID', 'Level', 'Notes']):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(10)
    for rid, desc, lvl, ts in rows:
        row = table.add_row().cells
        row[0].text = rid
        row[1].text = desc
        row[2].text = lvl
        row[3].text = ts
        for cell in row:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)


# ---------- TITLE ----------
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(p.add_run('SECURITY INCIDENT REPORT'), bold=True, size=20, color=NAVY)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(p.add_run('INC-2026-001: Unauthorised Local Account Creation and Privilege Escalation'), bold=True, size=12, color=NAVY)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(p.add_run('SOC Home Lab, Wazuh SIEM'), size=11, italic=True, color=GREY)

doc.add_paragraph()

# ---------- REPORT METADATA ----------
kv_table(doc, [
    ('Report Reference', 'INC-2026-001'),
    ('Date Compiled', '22 August 2026'),
    ('Analyst', 'Muhammad Ahmed Khalid'),
    ('System of Record', 'Wazuh SIEM (SOC Home Lab, 10.0.10.3)'),
    ('Classification', 'Internal, Training Exercise'),
    ('Status', 'Closed, post-incident review'),
    ('Overall Severity', 'Medium (highest Wazuh rule level observed: Level 12, Administrative Privilege Assigned)'),
    ('Simulation Reference', 'M5 Simulation 1, controlled attack simulation'),
])

# ---------- 1. EXECUTIVE SUMMARY ----------
h1(doc, '1. Executive Summary')
body(doc,
     'On 22 August 2026 at approximately 12:07 UTC, the Wazuh SIEM detected a suspicious sequence of Windows account management actions on the Windows-Endpoint host (10.0.10.4). A previously unknown local user account, "testuser", was created and immediately elevated to the local Administrators group. This activity is consistent with post-compromise persistence and privilege escalation techniques (MITRE ATT&CK T1136.001 and T1078.003) and would, in a live environment, require immediate investigation of the originating session, credentials, and any lateral movement.'
)
body(doc,
     'This incident was generated as part of a controlled training simulation within the SOC Home Lab to validate detection coverage and analyst workflow. The report is written as if the events were unplanned, so it can serve as a reference for future genuine incidents.'
)

# ---------- 2. INCIDENT DETAILS ----------
h1(doc, '2. Incident Details')
kv_table(doc, [
    ('Affected System', 'Windows-Endpoint (hostname WINDEV2407EVAL, IP 10.0.10.4)'),
    ('Affected Account', 'testuser (newly created local account)'),
    ('Time of Activity', '22 August 2026, approximately 12:07 UTC'),
    ('Time of Detection', '22 August 2026, approximately 12:07 UTC (within 60 seconds of activity)'),
    ('Detection Method', 'Wazuh SIEM correlation of Windows Security Event Log and Sysmon telemetry'),
    ('Attack Technique', 'Local account creation and administrative privilege escalation'),
    ('MITRE ATT&CK', 'T1136.001 (Create Account: Local Account), T1078.003 (Valid Accounts: Local Accounts)'),
    ('Total Wazuh Events Correlated', 'Approximately 20 events directly attributable to the activity, out of 463 events in the enclosing 24-hour window'),
])

# ---------- 3. TIMELINE ----------
h1(doc, '3. Timeline of Events')
body(doc,
     'The following timeline reconstructs the incident from the Wazuh alert view. All timestamps are approximate to the second; exact millisecond values are visible in the source screenshot in Appendix A.'
)
timeline_table(doc, [
    ('12:06:59', 'Attacker (simulated) session established: elevated PowerShell opened on Windows-Endpoint', 'Sysmon Event ID 1 (Process Create)'),
    ('12:06:59', 'Command executed: net user testuser P@ssw0rd123! /add', 'Windows Security Event 4720 + Sysmon'),
    ('12:07:00', 'Command executed: net localgroup administrators testuser /add', 'Windows Security Event 4732 + Sysmon'),
    ('12:07:00', 'Wazuh alert fired: Users Group Changed', 'Wazuh SIEM'),
    ('12:07:00', 'Wazuh alert fired: Administrative Privilege Assigned', 'Wazuh SIEM'),
    ('12:07:00', 'Wazuh alert fired: Domain Users Group Changed', 'Wazuh SIEM'),
    ('12:07:00', 'Wazuh alert fired: A net user account discovery command was executed', 'Wazuh SIEM'),
    ('12:07:00', 'Wazuh alert fired: Discovery activity spawned by PowerShell execution', 'Wazuh SIEM'),
    ('Approx. 12:08', 'Analyst identified the alert cluster during Threat Hunting review of Windows-Endpoint', 'Wazuh dashboard'),
    ('Approx. 12:10', 'Simulated containment: account marked for removal', 'Analyst action'),
])

# ---------- 4. DETECTION & ALERTING ----------
h1(doc, '4. Detection and Alerting')
body(doc,
     'The simulation triggered a cluster of correlated alerts across multiple Wazuh rules within the same second. Rule levels visible in the source screenshot ranged from 3 (informational) to 12 and above (high priority, would escalate in a live SOC). The most operationally significant alerts are listed first below.'
)
alert_table(doc, [
    ('Administrative Privilege Assigned', '60xxx range', '12', 'Direct hit from net localgroup command'),
    ('Users Group Changed', '60xxx range', '8', 'Direct hit from user creation'),
    ('Domain Users Group Changed', '60xxx range', '8', 'Related privilege change event'),
    ('A net user account discovery command was executed', '61606', '3', 'Detects the net.exe invocation'),
    ('Discovery activity spawned by PowerShell execution', '92xxx range', '5', 'Sysmon-based, catches parent-child process pattern'),
    ('System, Suspicious Process', '61151', '3', 'Multiple firings around net.exe and cmd.exe spawns'),
    ('Executable dropped in Windows directory', '61xxx range', '3', 'Related to process creation activity'),
])
body(doc,
     'Note: Exact rule IDs should be verified against the specific Wazuh ruleset version installed. Rule descriptions above are quoted directly from the Wazuh dashboard.',
     italic=True, color=GREY)

# ---------- 5. TECHNICAL ANALYSIS ----------
h1(doc, '5. Technical Analysis')
body(doc,
     'Two commands were executed on the Windows-Endpoint host from an elevated PowerShell session:'
)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
set_font(p.add_run('    net user testuser P@ssw0rd123! /add'), size=10)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(8)
set_font(p.add_run('    net localgroup administrators testuser /add'), size=10)

body(doc,
     'These commands together create a new local user account and grant it full administrative privileges on the host. Individually, each is a routine Windows administrator action. Executed together in rapid succession by an unexplained process, they represent a well-documented attacker post-compromise pattern:'
)
bullet(doc, 'T1136.001 Local Account, an attacker creates a local account for persistence, retaining access even if the original exploitation route is closed')
bullet(doc, 'T1078.003 Valid Accounts: Local Accounts, the attacker then uses the created account as if it were a legitimate user, evading behavioural detections that would otherwise flag anonymous activity')
body(doc,
     'Wazuh correlated Sysmon process-creation events (net.exe spawned by powershell.exe) with the Windows Security Event Log events (Event ID 4720 for account creation, Event ID 4732 for privileged group addition), and raised the corresponding alerts within seconds of the activity. The alert cluster is characteristic of this technique: several lower-level alerts appear alongside one or two high-severity alerts, and correlating them provides more confidence than any single alert alone.'
)

# ---------- 6. IMPACT ASSESSMENT ----------
h1(doc, '6. Impact Assessment')
kv_table(doc, [
    ('Confidentiality', 'Direct: none, no data was accessed. Potential (if unmitigated): full local access to any data readable by the Administrators group on the affected host.'),
    ('Integrity', 'Direct: unauthorised local account created and elevated. Potential (if unmitigated): complete host reconfiguration, log tampering, and persistent malware installation.'),
    ('Availability', 'Direct: none. Potential (if unmitigated): full host takeover, including denial of service to legitimate users.'),
    ('Blast Radius', 'Currently limited to the Windows-Endpoint host. Lateral movement paths from this host would need to be assessed in a live environment.'),
    ('Business Impact', 'Low (training environment). In a production equivalent, likely Medium to High depending on the affected host role.'),
])

# ---------- 7. RESPONSE ACTIONS ----------
h1(doc, '7. Response Actions')
h2(doc, 'Immediate, within first 5 minutes')
bullet(doc, 'Disable the newly created account: net user testuser /active:no')
bullet(doc, 'Force logout of any sessions belonging to testuser: quser then logoff [ID]')
bullet(doc, 'Isolate the host from the network (in a production SOC, via EDR isolation or firewall block)')
h2(doc, 'Short-term, within first hour')
bullet(doc, 'Identify the parent process and session that executed the net user commands using Sysmon Event ID 1 details (process image, command line, parent process, user context)')
bullet(doc, 'Search Wazuh for any additional accounts created within the same session window')
bullet(doc, 'Search for lateral movement indicators originating from this host (network connections to internal hosts, remote logins)')
h2(doc, 'Follow-up')
bullet(doc, 'Full removal of testuser account after evidence preservation: net user testuser /delete')
bullet(doc, 'Password reset for any accounts that were logged into the host during the incident window')
bullet(doc, 'Update endpoint hardening baseline if the initial access vector is identified')

# ---------- 8. ROOT CAUSE ----------
h1(doc, '8. Root Cause and Findings')
body(doc,
     'This simulation was executed by the analyst intentionally, so root cause in the technical sense is not applicable. The exercise did, however, validate the following about the SOC lab detection posture:'
)
bullet(doc, 'Windows Security Event Log ingestion via the Wazuh agent is working correctly, with expected alerts firing for account creation and privileged group changes')
bullet(doc, 'Sysmon telemetry with the SwiftOnSecurity baseline correctly identified the net.exe process spawn from powershell.exe as suspicious activity')
bullet(doc, 'Correlation between Sysmon process events and Windows Security events happened within the same alert window, indicating the Wazuh ruleset is matching effectively')
bullet(doc, 'Time from action to first visible alert was under 60 seconds, which meets Tier 1 SOC time-to-detect expectations for this class of activity')

# ---------- 9. RECOMMENDATIONS ----------
h1(doc, '9. Recommendations')
h2(doc, 'Detection Improvements')
bullet(doc, 'Enable real-time File Integrity Monitoring on high-value paths (see finding from M5 Simulation 4: a hosts file modification on Windows did not trigger a real-time FIM alert with the default configuration)')
bullet(doc, 'Add a custom correlation rule that fires a single high-severity alert when a new local account is added to Administrators within the same session, rather than relying on the analyst to correlate multiple lower-severity alerts')
h2(doc, 'Process Improvements')
bullet(doc, 'Document a standard analyst response runbook for the "unauthorised account creation" alert class')
bullet(doc, 'Baseline expected administrator activity per host, so anomalous privilege changes stand out more clearly')
h2(doc, 'Hardening Recommendations')
bullet(doc, 'Restrict local Administrators group membership to a documented list, with quarterly review')
bullet(doc, 'Enforce just-in-time privilege elevation for administrative tasks where practical')

# ---------- 10. LESSONS LEARNED ----------
h1(doc, '10. Lessons Learned')
bullet(doc, 'Multiple lower-severity alerts often tell a stronger story together than any single alert alone; analyst intuition and correlation matter as much as the tooling')
bullet(doc, 'Some Wazuh capabilities (like File Integrity Monitoring) are not enabled in real-time by default and require deliberate configuration for critical paths, this is a known gap to close in production')
bullet(doc, 'The 60-second detection window achieved here only holds because both endpoints are enrolled and telemetry is flowing; an unmanaged host would be blind, so onboarding discipline matters as much as detection engineering')

# ---------- APPENDIX A ----------
h1(doc, 'Appendix A, Evidence')
body(doc,
     'The following screenshots are stored in the project screenshots folder and are considered part of this report:'
)
bullet(doc, 'M5-sim1-user-created and user-added-to-admins.png, Wazuh Threat Hunting view for Windows-Endpoint immediately after the simulation. Shows the alert cluster including "Administrative Privilege Assigned", "Users Group Changed", "Domain Users Group Changed", "A net user account discovery command was executed", "Discovery activity spawned by PowerShell execution", and related Sysmon-based alerts.')

h1(doc, 'Appendix B, References')
bullet(doc, 'MITRE ATT&CK T1136.001, Create Account: Local Account, https://attack.mitre.org/techniques/T1136/001/')
bullet(doc, 'MITRE ATT&CK T1078.003, Valid Accounts: Local Accounts, https://attack.mitre.org/techniques/T1078/003/')
bullet(doc, 'Microsoft Windows Security Event 4720, A user account was created')
bullet(doc, 'Microsoft Windows Security Event 4732, A member was added to a security-enabled local group')
bullet(doc, 'Wazuh Documentation, https://documentation.wazuh.com')

# --------- FOOTER ---------
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(20)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(p.add_run('End of Report, INC-2026-001'), size=9, italic=True, color=GREY)

out = r'C:\Users\ahmed\0.CLAUDE\Job\Projects\Project 1 - SOC Home Lab with SIEM\incident-reports\INC-2026-001_Windows_Account_Creation.docx'
doc.save(out)
print(f'Saved: {out}')
